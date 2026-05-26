from pathlib import Path


_SUPPORTED = {".txt", ".md", ".csv", ".pdf", ".docx", ".xlsx"}
_MAX_BYTES = 300_000  # ~300 KB per file


def read_file(path_str: str) -> tuple[str, str]:
    """Read a local file and return (display_name, content).
    Raises ValueError with a user-friendly message on failure.
    """
    path = Path(path_str).expanduser().resolve()
    if not path.exists():
        raise ValueError(f"Arquivo não encontrado: {path}")
    if not path.is_file():
        raise ValueError(f"Não é um arquivo: {path}")

    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".xlsx":
        return _read_xlsx(path)

    if suffix not in {".txt", ".md", ".csv"}:
        raise ValueError(
            f"Formato não suportado: {suffix}. "
            "Use .pdf, .docx, .xlsx, .txt, .md ou .csv"
        )

    raw = path.read_bytes()
    if len(raw) > _MAX_BYTES:
        raw = raw[:_MAX_BYTES]
    content = raw.decode("utf-8", errors="replace")
    return path.name, content


def _read_pdf(path: Path) -> tuple[str, str]:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        raise ValueError("pypdf não instalado — execute: pip install pypdf")

    reader = PdfReader(str(path))
    pages: list[str] = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[Página {i + 1}]\n{text}")
    content = "\n\n".join(pages)
    if len(content.encode()) > _MAX_BYTES:
        content = content.encode()[:_MAX_BYTES].decode("utf-8", errors="replace")
    return path.name, content


def _read_docx(path: Path) -> tuple[str, str]:
    try:
        from docx import Document  # type: ignore
    except ImportError:
        raise ValueError("python-docx não instalado — execute: pip install python-docx")

    doc = Document(str(path))
    lines: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            lines.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    content = "\n".join(lines)
    if len(content.encode()) > _MAX_BYTES:
        content = content.encode()[:_MAX_BYTES].decode("utf-8", errors="replace")
    return path.name, content


def _read_xlsx(path: Path) -> tuple[str, str]:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        raise ValueError("openpyxl não instalado — execute: pip install openpyxl")

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"[Planilha: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                lines.append(" | ".join(cells))
    content = "\n".join(lines)
    if len(content.encode()) > _MAX_BYTES:
        content = content.encode()[:_MAX_BYTES].decode("utf-8", errors="replace")
    return path.name, content


def list_folder(path_str: str) -> list[Path]:
    """Return supported files in a folder (non-recursive)."""
    folder = Path(path_str).expanduser().resolve()
    if not folder.is_dir():
        raise ValueError(f"Pasta não encontrada: {folder}")
    return sorted(p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in _SUPPORTED)
